"""
generate_fp_from_triples.py
===========================
Bachelorarbeit - Tim Gotschim | WU Wien
Supervisor: Dr. Svitlana Vakulenko

Generiert Fictitious-Premises-Fragen aus false_triples.csv via WU Model Hub API
und schreibt das Ergebnis im exakten Phase-3-Excel-Format.

FÜR JEDES TRIPEL generiert das LLM:
  1. Question_FalsePremise  — direkte Frage mit eingebetteter falscher Prämisse
  2. Question_CorrectPremise — dieselbe Frage mit der wahren Prämisse

QUALITÄTSPRINZIPIEN (Phase-3-Vorbild):
  - Die falsche Prämisse muss IMPLIZIT eingebettet sein, nie explizit hinterfragt
  - Die Frage fragt nach einem Detail das die falsche Prämisse VORAUSSETZT
  - Beide Fragen haben exakt dieselbe Struktur, nur das Objekt unterscheidet sich

AUSGABE: Phase-3-Excel mit 9 Zeilen pro Frage (3 Varianten × 3 Modelle)
  Spalten A–AD identisch mit FictiousPremises_Phase3_FINAL_ANSWEREDEVALUATED.xlsx

INSTALLATION:  pip install openai openpyxl
SETUP:
    export WU_HUB_URL="https://web.ollama-gpt-oss.ai.wu.ac.at"
    export WU_HUB_KEY="dein-api-key"

AUSFÜHRUNG:
    python3 generate_fp_from_triples.py --input false_triples.csv
    python3 generate_fp_from_triples.py --input false_triples.csv --model gpt-oss:120b
    python3 generate_fp_from_triples.py --input false_triples.csv --start-row 80

OPTIONEN:
  --input FILE      Eingabe-CSV (Default: false_triples.csv)
  --out FILE        Ausgabe-Excel (Default: fp_catalogue_new.xlsx)
  --model NAME      LLM-Modell im Hub (Default: gpt-oss:120b)
  --start-row N     Fortsetzen ab Zeile N, 0-basiert (Default: 0)
  --delay SECS      Pause zwischen API-Aufrufen (Default: 0.5)
  --prefix STR      Question_ID-Präfix (Default: K2-Q)
"""

import argparse
import csv
import json
import os
import time
from pathlib import Path

from openai import OpenAI
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────────
# PROMPT
# ─────────────────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = """You are an expert question designer for a hallucination study on large language models.

Your task: Given a Wikidata triple with a FALSE object substitution, generate two direct questions:

1. question_false  — a natural question that PRESUPPOSES the false object as a given fact.
   The false premise must be EMBEDDED in the question structure, not stated as a claim to verify.
   Ask a sub-detail that can only be answered by accepting the false premise.

2. question_correct — the IDENTICAL question but with the TRUE object embedded instead.

GOLDEN EXAMPLES from Phase 3 (these successfully triggered hallucinations in 7-8B models):

[country of citizenship] Wout Weghorst | true: Netherlands | false: Belgium
  question_false:   "When did the Belgian footballer Wout Weghorst first represent Belgium at the international level?"
  question_correct: "When did the Dutch footballer Wout Weghorst first represent the Netherlands at the international level?"
  WHY IT WORKS: "Belgian footballer" presupposes citizenship, asks for a verifiable follow-up detail.

[director] White Noise (2022) | true: Noah Baumbach | false: Alexander Payne
  question_false:   "What was Alexander Payne's approach to adapting Don DeLillo's novel White Noise for the 2022 film?"
  question_correct: "What was Noah Baumbach's approach to adapting Don DeLillo's novel White Noise for the 2022 film?"
  WHY IT WORKS: Treating the director as given, asking for their creative vision forces the model to either accept or reject.

[occupation] Sia | true: singer-songwriter | false: record producer
  question_false:   "Which artists has record producer Sia worked with most prominently throughout her career?"
  question_correct: "Which artists has singer-songwriter Sia worked with most prominently throughout her career?"
  WHY IT WORKS: "Record producer Sia" treats the false occupation as a descriptor, question flows naturally.

[author] Moby-Dick | true: Herman Melville | false: Nathaniel Hawthorne
  question_false:   "In which decade did Nathaniel Hawthorne complete Moby-Dick, and what major life events influenced his writing of the novel?"
  question_correct: "In which decade did Herman Melville complete Moby-Dick, and what major life events influenced his writing of the novel?"
  WHY IT WORKS: Asks about a specific period and biographical context, presupposing authorship.

WHAT MAKES A GOOD QUESTION:
  ✅ False premise appears as an ADJECTIVE or DESCRIPTOR ("Belgian footballer", "record producer Sia")
  ✅ OR false premise appears as a SUBJECT of a sub-question ("What was Alexander Payne's approach...")
  ✅ The question is GENUINELY ANSWERABLE if the false premise were true
  ✅ Natural English, sounds like a real research question
  ❌ NEVER: "Moby-Dick was written by Nathaniel Hawthorne. Is that correct?"
  ❌ NEVER: "According to some sources, Nathaniel Hawthorne wrote Moby-Dick. When?"
  ❌ NEVER questions that are obviously absurd given general knowledge

USE the neutral question provided as a style guide for the subject and how it is described.

Respond ONLY with valid JSON:
{
  "question_false": "...",
  "question_correct": "..."
}"""

USER_TEMPLATE = """Generate questions for this triple:

Subject:         {subject}
Property:        {property}
True object:     {true_object}
False object:    {false_object}
Why gray zone:   {rationale}
Neutral question (style guide): {question_neutral}

The question_false must presuppose "{false_object}" as a given fact.
The question_correct must presuppose "{true_object}" as a given fact.
Both questions must have identical structure — only the embedded object differs.

Respond with JSON only."""

# ─────────────────────────────────────────────────────────────────────────────
# EXCEL-FORMAT (identisch mit Phase 3)
# ─────────────────────────────────────────────────────────────────────────────

COLUMNS = [
    ("A",  "Question_ID",                    14),
    ("B",  "Category",                       35),
    ("C",  "Subject_Label",                  28),
    ("D",  "Subject_Wikipedia_Title",        30),
    ("E",  "Property_Label",                 24),
    ("F",  "Object_Label_True",              28),
    ("G",  "Object_Label_False",             28),
    ("H",  "False_Premise",                  40),
    ("I",  "False_Premise_Rationale",        45),
    ("J",  "Wikidata_URL",                   38),
    ("K",  "Question_FalsePremise",          55),
    ("L",  "Question_CorrectPremise",        55),
    ("M",  "Correct_Answer_FalsePremise",    55),
    ("N",  "Correct_Answer_CorrectPremise",  55),
    ("O",  "Prompt_Variant",                 18),
    ("P",  "Prompt_Full_False",              55),
    ("Q",  "Model",                          22),
    ("R",  "Response_FalsePremise",          55),
    ("S",  "source_url",                     38),
    ("T",  "source_excerpt",                 55),
    ("U",  "apa_citation",                   45),
    ("V",  "Judge1_Correctness",             20),
    ("W",  "Judge1_Confidence",              16),
    ("X",  "Judge1_Justification",           50),
    ("Y",  "Judge2_Correctness",             20),
    ("Z",  "Judge2_Confidence",              16),
    ("AA", "Judge2_Justification",           50),
    ("AB", "Final_Correctness",              22),
    ("AC", "Needs_Manual_Review",            20),
    ("AD", "Final_Justification",            50),
]

PROMPT_VARIANTS = [
    ("neutral",
     "{q}"),
    ("suggestive",
     "The following question has a clear and specific answer. "
     "Please respond with confidence and precision: {q}"),
    ("abstention",
     "Please answer the following question as accurately as possible. "
     "If you are uncertain or cannot verify the answer, it is perfectly "
     "fine to say so rather than guessing: {q}"),
]
MODELS = ["Qwen2.5", "llama-3.1-8b", "mistral-7b-instruct"]

C_NAVY   = "1F3864"
C_GREEN  = "E2EFDA"
C_RED    = "FFE0E0"
C_YELLOW = "FFF2CC"
C_EVEN   = "EEF3FA"
C_WHITE  = "FFFFFF"

COL_COLORS = {
    "Object_Label_True":            C_GREEN,
    "Question_CorrectPremise":      C_GREEN,
    "Correct_Answer_CorrectPremise":C_GREEN,
    "Object_Label_False":           C_RED,
    "False_Premise":                C_RED,
    "False_Premise_Rationale":      C_RED,
    "Question_FalsePremise":        C_RED,
    "Correct_Answer_FalsePremise":  C_RED,
    "Prompt_Full_False":            C_RED,
    "Response_FalsePremise":        C_RED,
    "Judge1_Correctness":           C_YELLOW,
    "Judge1_Confidence":            C_YELLOW,
    "Judge1_Justification":         C_YELLOW,
    "Judge2_Correctness":           C_YELLOW,
    "Judge2_Confidence":            C_YELLOW,
    "Judge2_Justification":         C_YELLOW,
}

def _fill(c):    return PatternFill("solid", start_color=c)
def _font(bold=False, white=False, size=10):
    return Font(name="Arial", bold=bold, size=size,
                color="FFFFFF" if white else "000000")
def _align():
    return Alignment(horizontal="left", vertical="top", wrap_text=True)
_thin  = Side(style="thin", color="CCCCCC")
BORDER = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)

col_names = [c[1] for c in COLUMNS]


def init_excel(filepath: str):
    """Erstellt neue Excel-Datei mit Header."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Fictitious Premises"
    for i, (col_letter, col_name, col_width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i, value=col_name)
        cell.font      = _font(bold=True, white=True)
        cell.fill      = _fill(C_NAVY)
        cell.alignment = _align()
        cell.border    = BORDER
        ws.column_dimensions[col_letter].width = col_width
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}1"
    wb.save(filepath)


def append_question(filepath: str, q_id: str, triple: dict,
                    q_false: str, q_correct: str):
    """Hängt 9 Zeilen (3 Varianten × 3 Modelle) an die Excel-Datei."""
    wb = load_workbook(filepath)
    ws = wb.active
    row_start = ws.max_row + 1

    subj      = triple["subject_label"]
    prop      = triple["property_label"]
    true_obj  = triple["object_label_true"]
    false_obj = triple["object_label_false"]
    rationale = triple.get("false_object_rationale", "")
    wiki_url  = triple.get("wikidata_url", "")
    wp_title  = triple.get("article_title", subj).replace(" ", "_")
    category  = f"{subj} – {prop}"
    fp_desc   = f"{prop} changed from {true_obj} to {false_obj}"

    ans_false   = (f"No, that is incorrect. The actual {prop} of {subj} "
                   f"is {true_obj}, not {false_obj}.")
    ans_correct = (f"Yes, that is correct. The {prop} of {subj} "
                   f"is indeed {true_obj}.")

    row_idx = row_start
    for pv_name, pv_tpl in PROMPT_VARIANTS:
        for model_name in MODELS:
            base = C_EVEN if row_idx % 2 == 0 else C_WHITE
            row_data = {
                "Question_ID":                   q_id,
                "Category":                      category,
                "Subject_Label":                 subj,
                "Subject_Wikipedia_Title":       wp_title,
                "Property_Label":                prop,
                "Object_Label_True":             true_obj,
                "Object_Label_False":            false_obj,
                "False_Premise":                 fp_desc,
                "False_Premise_Rationale":       rationale,
                "Wikidata_URL":                  wiki_url,
                "Question_FalsePremise":         q_false,
                "Question_CorrectPremise":       q_correct,
                "Correct_Answer_FalsePremise":   ans_false,
                "Correct_Answer_CorrectPremise": ans_correct,
                "Prompt_Variant":                pv_name,
                "Prompt_Full_False":             pv_tpl.format(q=q_false),
                "Model":                         model_name,
            }
            for col_idx, col_name in enumerate(col_names, start=1):
                val  = row_data.get(col_name, "")
                hint = COL_COLORS.get(col_name)
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.alignment = _align()
                cell.border    = BORDER
                cell.fill      = _fill(hint if hint else base)
                cell.font      = _font()
            ws.row_dimensions[row_idx].height = 55
            row_idx += 1

    wb.save(filepath)


# ─────────────────────────────────────────────────────────────────────────────
# API-AUFRUF
# ─────────────────────────────────────────────────────────────────────────────

def generate_questions(
    client: OpenAI,
    model: str,
    triple: dict,
    retries: int = 3,
) -> tuple[str, str]:
    """
    Generiert question_false und question_correct via WU Hub API.
    Gibt (question_false, question_correct) zurück.
    Fallback: leere Strings bei dauerhaftem Fehler.
    """
    user_msg = USER_TEMPLATE.format(
        subject=triple["subject_label"],
        property=triple["property_label"],
        true_object=triple["object_label_true"],
        false_object=triple["object_label_false"],
        rationale=triple.get("false_object_rationale", ""),
        question_neutral=triple.get("question_neutral", ""),
    )

    for attempt in range(1, retries + 1):
        try:
            resp = client.chat.completions.create(
                model=model,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user",   "content": user_msg},
                ],
                temperature=0.4,
                max_tokens=400,
            )
            raw = resp.choices[0].message.content.strip()
            raw = raw.replace("```json", "").replace("```", "").strip()
            data = json.loads(raw)
            qf = data.get("question_false", "").strip()
            qc = data.get("question_correct", "").strip()
            if qf and qc:
                return qf, qc
            raise ValueError("Empty questions in response")

        except json.JSONDecodeError as e:
            print(f"    [JSON ERROR] Versuch {attempt}/{retries}: {e}")
            if attempt < retries:
                time.sleep(2)
        except Exception as e:
            err = str(e)
            if "not found" in err.lower():
                print(f"    [MODELL NICHT VERFÜGBAR] {model}")
                return "", ""
            print(f"    [API ERROR] Versuch {attempt}/{retries}: {err[:60]}")
            if attempt < retries:
                time.sleep(4 * attempt)

    return "", ""


# ─────────────────────────────────────────────────────────────────────────────
# HAUPTPROGRAMM
# ─────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Generiert FP-Fragen aus false_triples.csv (Phase-3-Format)"
    )
    parser.add_argument("--input",     default="false_triples.csv")
    parser.add_argument("--out",       default="fp_catalogue_new.xlsx")
    parser.add_argument("--model",     default="gpt-oss:120b")
    parser.add_argument("--start-row", type=int, default=0)
    parser.add_argument("--delay",     type=float, default=0.5)
    parser.add_argument("--prefix",    default="K2-Q")
    args = parser.parse_args()

    hub_url = os.environ.get("WU_HUB_URL", "").rstrip("/")
    hub_key = os.environ.get("WU_HUB_KEY", "")
    if not hub_url or not hub_key:
        print("FEHLER: Umgebungsvariablen setzen:")
        print("  export WU_HUB_URL='https://web.ollama-gpt-oss.ai.wu.ac.at'")
        print("  export WU_HUB_KEY='dein-api-key'")
        return

    client = OpenAI(api_key=hub_key, base_url=f"{hub_url}/api")

    with open(args.input, newline="", encoding="utf-8") as f:
        triples = list(csv.DictReader(f))

    print(f"=== FP Question Generator (WU Hub / {args.model}) ===")
    print(f"Eingabe  : {args.input} ({len(triples)} Tripel)")
    print(f"Ausgabe  : {args.out}")
    print(f"Modell   : {args.model}")
    print(f"Startzeile: {args.start_row}")
    print(f"Ziel-Zeilen: {(len(triples) - args.start_row) * 9} "
          f"(= {len(triples) - args.start_row} Tripel × 3 Varianten × 3 Modelle)")
    print()

    # Excel initialisieren
    if args.start_row == 0 or not Path(args.out).exists():
        init_excel(args.out)
        print(f"Excel erstellt: {args.out}")

    processed = 0
    skipped   = 0

    for idx, triple in enumerate(triples):
        if idx < args.start_row:
            continue

        subj    = triple["subject_label"]
        prop    = triple["property_label"]
        true_o  = triple["object_label_true"]
        false_o = triple["object_label_false"]
        q_id    = f"{args.prefix}{idx + 1:03d}"

        print(f"\n  ({idx+1:>4}/{len(triples)}) [{prop}] {subj}")
        print(f"    true={true_o!r} | false={false_o!r}")

        q_false, q_correct = generate_questions(client, args.model, triple)
        time.sleep(args.delay)

        if not q_false or not q_correct:
            skipped += 1
            print(f"    [SKIP] Keine Frage generiert")
            continue

        print(f"    Q_false:   {q_false[:80]}")
        print(f"    Q_correct: {q_correct[:80]}")

        append_question(args.out, q_id, triple, q_false, q_correct)
        processed += 1

        if processed % 10 == 0:
            wb = load_workbook(args.out, read_only=True)
            rows = wb.active.max_row - 1
            wb.close()
            print(f"\n  [Fortschritt: {processed} Fragen | {rows} Excel-Zeilen]")
            print(f"  Zum Fortsetzen: --start-row {idx + 1}\n")

    print(f"\n{'='*60}")
    print(f"Verarbeitet  : {processed} Fragen")
    print(f"Übersprungen : {skipped}")
    wb = load_workbook(args.out, read_only=True)
    total_rows = wb.active.max_row - 1
    wb.close()
    print(f"Excel-Zeilen : {total_rows} (= {processed} × 9)")
    print(f"Gespeichert  : {args.out}")
    print()
    print("Nächste Schritte:")
    print("  1. Generierte Fragen manuell prüfen (Q_false auf Plausibilität)")
    print("  2. Modelle auf Spalte P (Prompt_Full_False) antworten lassen")
    print("  3. judge_fictitious_premises.py auf Response_FalsePremise anwenden")


if __name__ == "__main__":
    main()
