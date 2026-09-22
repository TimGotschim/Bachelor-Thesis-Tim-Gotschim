"""
build_false_triples.py
======================
Bachelorarbeit - Tim Gotschim | WU Wien
Supervisor: Dr. Svitlana Vakulenko

Liest die verifizierten Tripel (alle drei Modelle beantworten die neutrale
Frage korrekt) und lässt Claude Sonnet für jedes Tripel ein verbessertes
falsches Objekt generieren, das den Phase-3-Qualitätskriterien entspricht.

INPUT:
  - neutral_questions_answered_evaluated.xlsx  (Knowledge_Check = Correct)
  - triples.csv                                (mit originalen falschen Objekten
                                                als Kontext/Ausgangspunkt)

OUTPUT:
  - false_triples.csv   (bereit für generate_fp_questions.py)

PHASE-3-QUALITÄTSKRITERIEN für das falsche Objekt ("Grauzone"):
  ✅ Gleiche semantische Klasse (Regisseur → anderer Regisseur, nie Politiker)
  ✅ Geografisch oder professionell benachbart (Portugal → Spanien)
  ✅ Plausibel auf den ersten Blick, aber definitiv falsch
  ✅ Gleiche Ära / Kulturkontext wie das wahre Objekt
  ❌ Nicht absurd (Stalin als Buchautor ist inakzeptabel)
  ❌ Nicht zu offensichtlich (USA → Kanada für US-Amerikaner)

INSTALLATION:  pip install anthropic openpyxl
SETUP:
    export ANTHROPIC_API_KEY="sk-ant-..."

AUSFÜHRUNG:
    python3 build_false_triples.py
    python3 build_false_triples.py --start-row 50   # Fortsetzen nach Abbruch
    python3 build_false_triples.py --input-eval neutral_questions_answered_evaluated.xlsx
                                   --input-triples triples.csv

OPTIONEN:
  --input-eval FILE     Evaluierte neutrale Fragen (Default: s.o.)
  --input-triples FILE  Tripel mit originalen falschen Objekten (Default: triples.csv)
  --out FILE            Ausgabe-CSV (Default: false_triples.csv)
  --start-row N         Fortsetzen ab Zeile N, 0-basiert (Default: 0)
  --delay SECS          Pause zwischen API-Aufrufen (Default: 0.5)
"""

import argparse
import csv
import json
import os
import time
from pathlib import Path

import anthropic
import openpyxl

# ─────────────────────────────────────────────────────────────────────────────
# CLAUDE SONNET PROMPT
# ─────────────────────────────────────────────────────────────────────────────

SYSTEM_PROMPT = """You are an expert at creating plausible-but-wrong alternatives for a hallucination
study on large language models (7–8B parameter models).

Your task: Given a verified fact triple (subject, property, true_object), generate a false_object
that belongs to the same semantic class as the true_object, but is definitively wrong.

The false_object must be in the "gray zone":
  ✅ Same type of entity as true_object (director → another real director, country → another country)
  ✅ Plausible at first glance due to geographic, professional, or cultural proximity
  ✅ Definitively wrong (not a matter of opinion or historical ambiguity)
  ✅ From the same era/context as the true_object
  ❌ Never absurd (no politicians as film directors, no athletes as book authors)
  ❌ Not too obvious (USA → Canada for an American is too easily dismissed)
  ❌ Not from a completely different context (Sanskrit text → Tatar language is wrong)

PROVEN EXAMPLES from Phase 3 that successfully triggered hallucinations:
  [country of citizenship] Wout Weghorst: Netherlands → Belgium
    WHY: Both football-mad nations, many players cross the border, plausible confusion.

  [director] White Noise (2022): Noah Baumbach → Alexander Payne
    WHY: Both acclaimed US indie directors known for literary adaptations.

  [occupation] Sia: singer-songwriter → record producer
    WHY: Sia genuinely produces music for others, roles overlap in music industry.

  [country of citizenship] João Félix: Portugal → Spain
    WHY: Plays for Spanish clubs, Iberian peninsula, frequent transfers.

  [language of work] Don Quixote: Early Modern Spanish → Portuguese
    WHY: Both Iberian languages, same era, closely related.

  [author] A Christmas Carol: Charles Dickens → William Makepeace Thackeray
    WHY: Both Victorian British novelists, contemporaries, same social realism style.

  [director] Skyfall: Sam Mendes → Christopher Nolan
    WHY: Both acclaimed British directors of prestige Hollywood films in same era.

  [country of origin] The Hobbit: United Kingdom → New Zealand
    WHY: The film was actually shot in New Zealand, the confusion is very natural.

EXAMPLES OF BAD false objects (do NOT do this):
  ❌ [director] Thor: Kenneth Branagh → Vladimir Lenin (not a director)
  ❌ [author] Pride and Prejudice: Jane Austen → Greg Rutherford (an athlete)
  ❌ [language of work] Moby-Dick: English → Tatar (no connection whatsoever)
  ❌ [occupation] Garry Kasparov: chess player → sniper (absurd)
  ❌ [country of citizenship] Garry Kasparov: Soviet Union → Achaemenid Empire (wrong era)

Respond ONLY with valid JSON, no other text:
{
  "false_object": "the plausible but wrong alternative",
  "rationale": "one sentence: why is this in the gray zone?"
}"""

USER_TEMPLATE = """Generate a false_object for this triple:

Subject:      {subject}
Property:     {property}
True object:  {true_object}
Original false object (CURRENT, often bad quality): {original_false}
Neutral question asked about this subject: {neutral_question}

The false_object must be of the same type as "{true_object}" but definitively wrong.
Use the neutral question for context about how the subject is described.

Respond with JSON only."""

# ─────────────────────────────────────────────────────────────────────────────
# OUTPUT-FELDER
# ─────────────────────────────────────────────────────────────────────────────

OUTPUT_FIELDS = [
    "article_title",
    "subject_label",
    "property_label",
    "object_label_true",
    "object_label_false",          # verbessertes falsches Objekt
    "object_label_false_original", # originales (schlechtes) falsches Objekt
    "false_object_rationale",      # Begründung von Claude Sonnet
    "question_neutral",            # neutrale Frage (für Kontext)
    "wikidata_url",
    "wikipedia_url",
    "avg_views_2023",
]

# ─────────────────────────────────────────────────────────────────────────────
# API-AUFRUF
# ─────────────────────────────────────────────────────────────────────────────

def get_false_object(
    client: anthropic.Anthropic,
    subject: str,
    property_label: str,
    true_object: str,
    original_false: str,
    neutral_question: str,
    retries: int = 3,
) -> dict:
    """
    Fragt Claude Sonnet nach einem verbesserten falschen Objekt.
    Gibt {'false_object': ..., 'rationale': ...} zurück.
    Bei Fehler: originales falsches Objekt als Fallback.
    """
    user_msg = USER_TEMPLATE.format(
        subject=subject,
        property=property_label,
        true_object=true_object,
        original_false=original_false or "(none)",
        neutral_question=neutral_question or "(not available)",
    )

    for attempt in range(1, retries + 1):
        try:
            response = client.messages.create(
                model="claude-sonnet-4-20250514",
                max_tokens=200,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": user_msg}],
            )
            raw = response.content[0].text.strip()
            raw = raw.replace("```json", "").replace("```", "").strip()
            result = json.loads(raw)

            fo = result.get("false_object", "").strip()
            rationale = result.get("rationale", "").strip()

            # Validierung: falsches Objekt darf nicht gleich dem wahren sein
            if fo and fo.lower() != true_object.lower():
                return {"false_object": fo, "rationale": rationale}

            return {"false_object": original_false, "rationale": "fallback (validation failed)"}

        except json.JSONDecodeError:
            if attempt < retries:
                time.sleep(2)
        except Exception as e:
            print(f"    [API ERROR] Versuch {attempt}/{retries}: {str(e)[:60]}")
            if attempt < retries:
                time.sleep(4 * attempt)

    # Fallback: originales falsches Objekt behalten
    return {"false_object": original_false, "rationale": "fallback (max retries)"}

# ─────────────────────────────────────────────────────────────────────────────
# HAUPTPROGRAMM
# ─────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Generiert verbesserte falsche Tripel via Claude Sonnet"
    )
    parser.add_argument("--input-eval",
        default="neutral_questions_answered_evaluated.xlsx")
    parser.add_argument("--input-triples",
        default="triples.csv")
    parser.add_argument("--out",
        default="false_triples.csv")
    parser.add_argument("--start-row", type=int, default=0,
        help="0-basiert; zum Fortsetzen nach Abbruch")
    parser.add_argument("--delay",     type=float, default=0.5)
    args = parser.parse_args()

    api_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if not api_key:
        print("FEHLER: ANTHROPIC_API_KEY nicht gesetzt.")
        print("  export ANTHROPIC_API_KEY='sk-ant-...'")
        return

    client = anthropic.Anthropic(api_key=api_key)

    # ── Triples-Lookup laden ─────────────────────────────────────────────────
    with open(args.input_triples, newline="", encoding="utf-8") as f:
        triples_lookup = {
            (r["article_title"], r["property_label"]): r
            for r in csv.DictReader(f)
        }

    # ── Evaluierte neutrale Fragen laden (nur Correct) ──────────────────────
    wb = openpyxl.load_workbook(args.input_eval)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    def col(name):
        return headers.index(name) + 1

    correct_rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col("Knowledge_Check")).value != "Correct":
            continue
        correct_rows.append({
            "article_title":    ws.cell(r, col("article_title")).value  or "",
            "subject_label":    ws.cell(r, col("subject_label")).value  or "",
            "property_label":   ws.cell(r, col("property_label")).value or "",
            "object_label_true":ws.cell(r, col("object_label_true")).value or "",
            "question_neutral": ws.cell(r, col("Question")).value       or "",
        })

    total = len(correct_rows)
    print(f"=== False Triple Builder (Claude Sonnet) ===")
    print(f"Eingabe:    {args.input_eval} → {total} Correct-Tripel")
    print(f"Ausgabe:    {args.out}")
    print(f"Startzeile: {args.start_row}")
    print()

    # ── Bestehende Ausgabe laden (Fortsetzen) ────────────────────────────────
    existing_results = []
    if args.start_row > 0 and Path(args.out).exists():
        with open(args.out, newline="", encoding="utf-8") as f:
            existing_results = list(csv.DictReader(f))
        print(f"Fortsetzen: {len(existing_results)} bereits verarbeitet")

    # ── Ausgabedatei öffnen ──────────────────────────────────────────────────
    write_header = args.start_row == 0 or not Path(args.out).exists()
    out_file = open(args.out,
                    "w" if write_header else "a",
                    newline="", encoding="utf-8")
    writer = csv.DictWriter(out_file, fieldnames=OUTPUT_FIELDS)
    if write_header:
        writer.writeheader()

    # ── Verarbeitung ─────────────────────────────────────────────────────────
    processed = 0
    improved  = 0

    for idx, row in enumerate(correct_rows):
        if idx < args.start_row:
            continue

        subject    = row["subject_label"]
        prop       = row["property_label"]
        true_obj   = row["object_label_true"]
        q_neutral  = row["question_neutral"]
        art_title  = row["article_title"]

        # Originales falsches Objekt aus triples.csv holen
        triple_key = (art_title, prop)
        original   = triples_lookup.get(triple_key, {})
        orig_false = original.get("object_label_false", "")
        wikidata_url   = original.get("wikidata_url", "")
        wikipedia_url  = original.get("wikipedia_url", "")
        avg_views      = original.get("avg_views_2023", "")

        print(f"  ({idx+1:>4}/{total}) [{prop}] {subject}")
        print(f"    true:     {true_obj}")
        print(f"    orig:     {orig_false}")

        result = get_false_object(
            client, subject, prop, true_obj, orig_false, q_neutral
        )
        time.sleep(args.delay)

        new_false = result["false_object"]
        rationale = result["rationale"]
        was_improved = new_false != orig_false

        if was_improved:
            improved += 1
            print(f"    improved: {new_false}  ✅")
        else:
            print(f"    kept:     {new_false}")
        print(f"    why:      {rationale}")

        writer.writerow({
            "article_title":            art_title,
            "subject_label":            subject,
            "property_label":           prop,
            "object_label_true":        true_obj,
            "object_label_false":       new_false,
            "object_label_false_original": orig_false,
            "false_object_rationale":   rationale,
            "question_neutral":         q_neutral,
            "wikidata_url":             wikidata_url,
            "wikipedia_url":            wikipedia_url,
            "avg_views_2023":           avg_views,
        })
        out_file.flush()
        processed += 1

        if processed % 20 == 0:
            print(f"\n  [Auto-Save: {processed} verarbeitet | "
                  f"Fortsetzen mit --start-row {idx + 1}]\n")

    out_file.close()

    print(f"\n{'='*55}")
    print(f"Verarbeitet  : {processed}")
    print(f"Verbessert   : {improved} ({improved/max(processed,1)*100:.0f}%)")
    print(f"Unverändert  : {processed - improved}")
    print(f"Gespeichert  : {args.out}")
    print()
    print("Nächster Schritt:")
    print("  python3 generate_fp_questions.py --input false_triples.csv")
    print("  → generiert Fragen mit falscher Prämisse im Phase-3-Format")


if __name__ == "__main__":
    main()
