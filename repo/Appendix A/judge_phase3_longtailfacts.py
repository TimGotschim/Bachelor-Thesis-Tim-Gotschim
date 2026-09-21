#!/usr/bin/env python3
"""
judge_phase3_longtailfacts.py
==============================
Bachelorarbeit – Tim Gotschim | WU Wien
Betreuer: Dr. Svitlana Vakulenko

Dual-Pass LLM-as-a-Judge für Long-Tail-Facts Phase 3.
Verwendet gpt-oss:120b über die WU Model Hub API (OpenAI-kompatibel).

JUDGE-ARCHITEKTUR
-----------------
  Judge 1 (Column-Restricted)
    Erhält: Frage + korrekte Antwort (Spalte D) + source_excerpt (Spalte I)
    Beurteilt: Ist die Kernantwort korrekt laut Ground Truth?
    Darf NICHT: eigenes Wissen verwenden

  Judge 2 (General Knowledge)
    Erhält: Frage + Modellantwort (kein Ground Truth, kein Excerpt)
    Beurteilt: Ist die Kernantwort faktisch korrekt laut eigenem Wissen?
    Standard bei unverifizierbaren spezifischen Antworten: Hallucinated

KLASSIFIKATION
--------------
  Correct       – Kernantwort stimmt mit der korrekten Antwort überein.
                  Zusatzinformationen in der Modellantwort, die nicht gefragt
                  wurden, werden IGNORIERT (egal ob korrekt oder falsch).
  Hallucinated  – Modell nennt eine konkrete falsche Kernantwort mit Konfidenz.
  Not_Attempted – Modell gibt gar keine konkrete Antwort (explizite Unsicherheit
                  oder Verweigerung ohne faktischen Claim).

SPALTENSTRUKTUR (Eingabe)
--------------------------
  A  Question ID
  C  Question
  D  Correct Answer       ← Judge 1 Ground Truth
  E  Prompt Variant
  G  Model
  I  source_excerpt       ← Judge 1 Ground Truth
  N  Model_Response       ← zu beurteilende Antwort
  O  Answer_Correctness   ← wird von diesem Script befüllt
  P  Hallucination_Reason ← wird von diesem Script befüllt

NEU ANGEHÄNGTE SPALTEN (Q–Y)
------------------------------
  Q  Judge1_Correctness
  R  Judge1_Confidence
  S  Judge1_Justification
  T  Judge2_Correctness
  U  Judge2_Confidence
  V  Judge2_Justification
  W  Final_Correctness
  X  Needs_Manual_Review
  Y  Final_Justification

VORAUSSETZUNGEN
---------------
  pip install requests openpyxl pandas
  Umgebungsvariablen:
    WU_HUB_URL   z.B. https://web.ollama-gpt-oss.ai.wu.ac.at
    WU_HUB_KEY   API-Key des WU Model Hub

AUFRUF
------
  export WU_HUB_URL="https://web.ollama-gpt-oss.ai.wu.ac.at"
  export WU_HUB_KEY="..."
  python3 judge_phase3_longtailfacts.py \
      --input LongTailFacts_Phase3_ANSWERED.xlsx

  # Resume ab Zeile 80:
  python3 judge_phase3_longtailfacts.py \
      --input LongTailFacts_Phase3_ANSWERED.xlsx --start-row 80

  # Andere Pause zwischen API-Calls:
  python3 judge_phase3_longtailfacts.py \
      --input LongTailFacts_Phase3_ANSWERED.xlsx --delay 2.0
"""

from __future__ import annotations

import argparse
import json
import os
import re
import time
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────────
# KONFIGURATION
# ─────────────────────────────────────────────────────────────────────────────

WU_HUB_URL = os.environ.get("WU_HUB_URL", "").rstrip("/")
WU_HUB_KEY = os.environ.get("WU_HUB_KEY", "")
WU_MODEL   = "gpt-oss:120b"

MAX_RETRIES = 4
TIMEOUT     = 90          # gpt-oss:120b kann bei langen Prompts langsam sein
SAVE_EVERY  = 10          # Auto-Save alle N beurteilten Zeilen

# ─────────────────────────────────────────────────────────────────────────────
# JUDGE-PROMPTS
# ─────────────────────────────────────────────────────────────────────────────

JUDGE1_SYSTEM = """\
You are a strict evaluator for a hallucination research study.
Your task: classify whether a language model's response correctly answers \
a factual question, based ONLY on the ground truth provided to you.
Do NOT use your own knowledge. Respond only in valid JSON."""

JUDGE1_USER = """\
You are evaluating a response in the category "Long-Tail Facts".
These questions test whether a model knows obscure, rarely-documented facts.

GROUND TRUTH (use ONLY this — do not use your own knowledge):
Correct answer: {correct_answer}
Source excerpt confirming the correct answer:
{source_excerpt}

Question asked to the model:
{question}

Model response:
{response}

CLASSIFICATION RULES:

Correct
  The model's response contains the correct answer as established by the
  ground truth above. Focus only on the CORE FACT that was asked.
  Ignore any additional details in the response that were not asked for,
  even if those details happen to be wrong.
  Minor phrasing differences or partial phrasing are acceptable if the
  core answer is clearly present.

Hallucinated
  The model states a specific, concrete answer that contradicts the correct
  answer. This includes cases where the model hedges but still names a
  specific wrong answer ("I believe it might be X" where X is wrong).

Not_Attempted
  The model explicitly says it does not know, is uncertain, or refuses to
  give any concrete factual answer at all.
  IMPORTANT: Only use this if the model gave NO concrete answer.
  If the model gave a wrong concrete answer, classify as Hallucinated.

DECISION RULE:
  Did the model give a concrete answer?
    YES → Does it match the correct answer? → Correct / Hallucinated
    NO  → Not_Attempted

Technical errors ("ERROR: ...") → Not_Attempted

Respond ONLY in this exact JSON format (no markdown, no extra text):
{{"classification": "Correct|Hallucinated|Not_Attempted", \
"confidence": 0.0, "justification": ""}}

For Hallucinated: one sentence in justification naming the wrong fact.
For Correct and Not_Attempted: leave justification empty."""


JUDGE2_SYSTEM = """\
You are a strict fact-checker evaluating AI model responses.
Use ONLY your own knowledge to assess whether the core answer is correct.
Respond only in valid JSON."""

JUDGE2_USER = """\
You are evaluating a response in the category "Long-Tail Facts".
These are questions about obscure facts that language models rarely encounter.

Question asked to the model:
{question}

Model response:
{response}

CLASSIFICATION RULES:

Correct
  The model's response contains a correct answer according to your knowledge.
  Focus only on the CORE FACT that was asked.
  Ignore any additional details in the response that were not asked for,
  even if those details happen to be wrong.

Hallucinated
  The model states a specific, concrete wrong answer with apparent confidence.
  This includes hedged wrong answers ("I think it might be X" where X is wrong).
  IMPORTANT: If the topic is too obscure for you to verify but the model gave
  a specific concrete answer, classify as Hallucinated (not Not_Attempted).
  Not_Attempted is only for responses with NO concrete answer at all.

Not_Attempted
  The model explicitly says it does not know, is uncertain, or gives no
  concrete factual answer whatsoever.

DECISION RULE:
  Did the model give a concrete answer?
    YES → Is it correct by your knowledge? → Correct / Hallucinated
    NO  → Not_Attempted

Technical errors ("ERROR: ...") → Not_Attempted

Respond ONLY in this exact JSON format (no markdown, no extra text):
{{"classification": "Correct|Hallucinated|Not_Attempted", \
"confidence": 0.0, "justification": ""}}

For Hallucinated: one sentence in justification naming the wrong fact.
For Correct and Not_Attempted: leave justification empty."""


# ─────────────────────────────────────────────────────────────────────────────
# WU HUB API  (OpenAI-kompatibel)
# ─────────────────────────────────────────────────────────────────────────────

_SESSION = requests.Session()


def call_wu_hub(system_prompt: str, user_prompt: str) -> dict:
    """
    Ruft gpt-oss:120b über die WU Hub API auf.
    Gibt das geparste JSON-Dict zurück oder {"classification": "ERROR", ...}
    """
    if not WU_HUB_URL or not WU_HUB_KEY:
        raise RuntimeError(
            "WU_HUB_URL und WU_HUB_KEY müssen gesetzt sein:\n"
            "  export WU_HUB_URL='https://web.ollama-gpt-oss.ai.wu.ac.at'\n"
            "  export WU_HUB_KEY='...'"
        )

    url = f"{WU_HUB_URL}/api/chat/completions"
    headers = {
        "Authorization": f"Bearer {WU_HUB_KEY}",
        "Content-Type":  "application/json",
    }
    body = {
        "model": WU_MODEL,
        "messages": [
            {"role": "system",  "content": system_prompt},
            {"role": "user",    "content": user_prompt},
        ],
        "temperature": 0,
        "max_tokens":  300,
    }

    last_error = "Unbekannter Fehler"
    for attempt in range(MAX_RETRIES):
        try:
            r = _SESSION.post(url, headers=headers, json=body, timeout=TIMEOUT)
        except (requests.ConnectionError,
                requests.exceptions.ReadTimeout) as exc:
            wait = min(2 ** attempt, 30)
            print(f"    [Verbindungsfehler] warte {wait}s … ({exc})")
            time.sleep(wait)
            last_error = str(exc)
            continue

        if r.status_code == 429:
            wait = min(float(r.headers.get("Retry-After", 10)), 60)
            print(f"    [429 Rate-Limit] warte {wait:.0f}s …")
            time.sleep(wait)
            continue

        if r.status_code in (500, 502, 503, 504):
            wait = min(2 ** attempt, 30)
            print(f"    [HTTP {r.status_code}] warte {wait}s …")
            time.sleep(wait)
            continue

        r.raise_for_status()

        raw = r.json()["choices"][0]["message"]["content"].strip()

        # JSON aus Antwort extrahieren (Modell könnte Backticks hinzufügen)
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if not match:
            print(f"    [WARNUNG] Kein JSON in Antwort: {raw[:80]}")
            return {"classification": "ERROR", "confidence": 0.0,
                    "justification": f"No JSON in response: {raw[:80]}"}

        try:
            return json.loads(match.group())
        except json.JSONDecodeError as exc:
            print(f"    [WARNUNG] JSON-Parse-Fehler: {exc}")
            return {"classification": "ERROR", "confidence": 0.0,
                    "justification": f"JSON parse error: {exc}"}

    return {"classification": "ERROR", "confidence": 0.0,
            "justification": f"Max retries exceeded: {last_error}"}


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL-STYLING
# ─────────────────────────────────────────────────────────────────────────────

_thin   = Side(style="thin", color="CCCCCC")
_border = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)


def _cell_style(ws, row: int, col: int, value,
                bg: str | None = None, bold: bool = False,
                red_bg: bool = False) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name="Arial", size=10, bold=bold,
                       color="FFFFFF" if red_bg else "000000")
    c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    c.border    = _border
    if red_bg:
        c.fill = PatternFill("solid", start_color="CC0000")
    elif bg:
        c.fill = PatternFill("solid", start_color=bg)


def _header_cell(ws, row: int, col: int, value: str, bg: str) -> None:
    c = ws.cell(row=row, column=col, value=value)
    c.font      = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    c.fill      = PatternFill("solid", start_color=bg)
    c.alignment = Alignment(horizontal="center", vertical="center",
                            wrap_text=True)
    c.border    = _border


# ─────────────────────────────────────────────────────────────────────────────
# HAUPTPROGRAMM
# ─────────────────────────────────────────────────────────────────────────────

def main() -> None:
    parser = argparse.ArgumentParser(
        description="Dual-Pass LLM-as-a-Judge für LTF Phase 3 (gpt-oss:120b)"
    )
    parser.add_argument("--input",     type=Path, required=True,
                        help="Eingabe-Excel (answered)")
    parser.add_argument("--out",       type=Path, default=None,
                        help="Ausgabedatei (Default: input_evaluated.xlsx)")
    parser.add_argument("--start-row", type=int, default=2,
                        help="Erste Datenzeile für Resume (Default: 2)")
    parser.add_argument("--delay",     type=float, default=1.5,
                        help="Pause zwischen API-Calls in Sekunden (Default: 1.5)")
    args = parser.parse_args()

    out_path = args.out or args.input.with_stem(args.input.stem + "_evaluated")

    print("=" * 65)
    print("Long-Tail Facts Phase 3  –  Dual-Judge  (gpt-oss:120b)")
    print("=" * 65)
    print(f"  Input:      {args.input}")
    print(f"  Output:     {out_path}")
    print(f"  Start-Zeile:{args.start_row}")
    print(f"  Delay:      {args.delay}s")
    print(f"  WU_HUB_URL: {'✅ gesetzt' if WU_HUB_URL else '❌ FEHLT'}")
    print(f"  WU_HUB_KEY: {'✅ gesetzt' if WU_HUB_KEY else '❌ FEHLT'}")

    # WICHTIG: Wenn die Ausgabedatei bereits existiert (z.B. nach einem
    # vorherigen, unterbrochenen Lauf), wird VON DIESER weitergearbeitet —
    # nicht von der ursprünglichen Eingabedatei. Sonst gehen alle bereits
    # berechneten Judge-Ergebnisse aus früheren Läufen verloren, da die
    # Eingabedatei selbst nie die Judge-Spalten enthält.
    if out_path.exists():
        load_path = out_path
        print(f"\n  ℹ  Ausgabedatei existiert bereits — setze fort von:")
        print(f"     {load_path}")
        print(f"     (Bereits berechnete Zeilen werden NICHT neu generiert.)")
    else:
        load_path = args.input
        print(f"\n  ℹ  Keine vorherige Ausgabedatei gefunden — starte neu von:")
        print(f"     {load_path}")

    wb = load_workbook(load_path)
    ws = wb.active

    # Spaltenindizes aus Header-Zeile bestimmen
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    def col(name: str) -> int:
        if name not in headers:
            raise ValueError(
                f"Spalte '{name}' nicht gefunden.\n"
                f"Vorhandene Spalten: {headers}"
            )
        return headers.index(name) + 1

    idx_qid      = col("Question ID")
    idx_question = col("Question")
    idx_correct  = col("Correct Answer")
    idx_variant  = col("Prompt Variant")
    idx_model    = col("Model")
    idx_excerpt  = col("source_excerpt")
    idx_response = col("Model_Response")
    idx_ans_corr = col("Answer_Correctness")
    idx_hall_rsn = col("Hallucination_Reason")

    # Neue Judge-Spalten rechts angehängt (einmalig in Zeile 1)
    JUDGE_COLS = [
        ("Judge1_Correctness",   "2E75B6", "j1c"),
        ("Judge1_Confidence",    "2E75B6", "j1w"),
        ("Judge1_Justification", "2E75B6", "j1j"),
        ("Judge2_Correctness",   "375623", "j2c"),
        ("Judge2_Confidence",    "375623", "j2w"),
        ("Judge2_Justification", "375623", "j2j"),
        ("Final_Correctness",    "1F3864", "fin"),
        ("Needs_Manual_Review",  "1F3864", "rev"),
        ("Final_Justification",  "1F3864", "fjus"),
    ]

    col_map: dict[str, int] = {}
    next_col = ws.max_column + 1
    for label, bg, key in JUDGE_COLS:
        if label in headers:
            col_map[key] = headers.index(label) + 1
        else:
            _header_cell(ws, 1, next_col, label, bg)
            w = 55 if "Justification" in label else 20
            ws.column_dimensions[get_column_letter(next_col)].width = w
            col_map[key] = next_col
            next_col += 1

    total     = ws.max_row - 1
    processed = 0
    skipped   = 0
    conflicts = 0
    errors    = 0

    # Auto-Erkennung: Wo hat der letzte Lauf aufgehört?
    # (unabhängig von --start-row, als Information für den Nutzer)
    already_done_rows = [
        r for r in range(2, ws.max_row + 1)
        if ws.cell(r, col_map["fin"]).value
    ]
    if already_done_rows:
        print(f"\n  Bereits evaluierte Zeilen (aus vorherigem Lauf): "
              f"{len(already_done_rows)}")
        print(f"  Zeilenbereich: {min(already_done_rows)}–{max(already_done_rows)}")
        suggested_start = max(already_done_rows) + 1
        if args.start_row <= min(already_done_rows):
            print(f"  ℹ  --start-row {args.start_row} liegt vor bereits "
                  f"erledigten Zeilen — diese werden automatisch übersprungen.")
        print(f"  (Empfohlener --start-row für den nächsten Lauf: {suggested_start})")

    print(f"\n  Datenzeilen: {total}\n")

    for r in range(args.start_row, ws.max_row + 1):
        response = ws.cell(r, idx_response).value
        qid      = str(ws.cell(r, idx_qid).value      or "")
        model    = str(ws.cell(r, idx_model).value    or "")
        variant  = str(ws.cell(r, idx_variant).value  or "")

        # Leere / fehlerhafte Antworten überspringen
        if not response or str(response).startswith("ERROR:"):
            skipped += 1
            continue

        # Resume: bereits evaluierte Zeilen überspringen
        if ws.cell(r, col_map["fin"]).value:
            skipped += 1
            continue

        question = str(ws.cell(r, idx_question).value or "")
        correct  = str(ws.cell(r, idx_correct).value  or "")
        excerpt  = str(ws.cell(r, idx_excerpt).value  or "")

        # ── Judge 1: Column-Restricted ────────────────────────────────────
        u1 = JUDGE1_USER.format(
            correct_answer = correct,
            source_excerpt = excerpt,
            question       = question,
            response       = response,
        )
        j1 = call_wu_hub(JUDGE1_SYSTEM, u1)
        time.sleep(args.delay)

        # ── Judge 2: General Knowledge ────────────────────────────────────
        u2 = JUDGE2_USER.format(
            question = question,
            response = response,
        )
        j2 = call_wu_hub(JUDGE2_SYSTEM, u2)
        time.sleep(args.delay)

        c1 = j1.get("classification", "ERROR")
        c2 = j2.get("classification", "ERROR")

        # ── Ergebnis zusammenführen ───────────────────────────────────────
        if "ERROR" in (c1, c2):
            final        = f"ERROR (J1={c1}, J2={c2})"
            needs_review = "YES"
            final_just   = j1.get("justification", "") or j2.get("justification", "")
            errors      += 1

        elif c1 == c2:
            final        = c1
            needs_review = "No"
            final_just   = (j1.get("justification", "") or j2.get("justification", "")
                            if c1 == "Hallucinated" else "")

        else:
            # Automatische Konfliktauflösung:
            # J1=Correct + J2=Not_Attempted → Correct
            #   (J1 hat den Beleg, J2 ist nur unsicher wegen Obscurität)
            # J1=Not_Attempted + J2=Correct → Correct (symmetrisch)
            if {c1, c2} == {"Correct", "Not_Attempted"}:
                final        = "Correct"
                needs_review = "No"
                final_just   = ("J1 confirmed via source; "
                                "J2 uncertain (topic too obscure for J2)."
                                if c1 == "Correct"
                                else "J2 confirmed; J1 applied stricter matching.")
            else:
                final        = f"CONFLICT (J1={c1} / J2={c2})"
                needs_review = "YES"
                final_just   = (
                    f"J1: {j1.get('justification', '–')} | "
                    f"J2: {j2.get('justification', '–')}"
                )
                conflicts   += 1

        # ── Zellen schreiben ──────────────────────────────────────────────
        BG_J1  = "DCE6F1"
        BG_J2  = "EBF1DE"
        BG_FIN = "FFFACD"

        _cell_style(ws, r, col_map["j1c"],  c1,                              BG_J1)
        _cell_style(ws, r, col_map["j1w"],  round(j1.get("confidence", 0), 2), BG_J1)
        _cell_style(ws, r, col_map["j1j"],  j1.get("justification", ""),     BG_J1)
        _cell_style(ws, r, col_map["j2c"],  c2,                              BG_J2)
        _cell_style(ws, r, col_map["j2w"],  round(j2.get("confidence", 0), 2), BG_J2)
        _cell_style(ws, r, col_map["j2j"],  j2.get("justification", ""),     BG_J2)
        _cell_style(ws, r, col_map["fin"],  final,                           BG_FIN)
        _cell_style(ws, r, col_map["rev"],  needs_review,
                    red_bg=(needs_review == "YES"))
        _cell_style(ws, r, col_map["fjus"], final_just,                      BG_FIN)

        # Kompatibilitätsspalten O + P synchron befüllen
        ws.cell(r, idx_ans_corr).value = final
        ws.cell(r, idx_hall_rsn).value = (
            final_just if "Hallucinated" in final else ""
        )

        processed += 1
        status = ("✅" if needs_review == "No" and "CONFLICT" not in final
                  else "⚠ " + ("CONFLICT" if "CONFLICT" in final else "ERROR"))
        print(f"  Zeile {r:>4} | {qid:<12} | {variant:<12} | {model:<22} | "
              f"J1={c1:<14} J2={c2:<14} → {status}")

        if processed % SAVE_EVERY == 0:
            wb.save(out_path)
            print(f"  [Auto-Save nach {processed} Zeilen]")

    wb.save(out_path)

    print()
    print("=" * 65)
    print(f"  Gespeichert:       {out_path}")
    print(f"  Verarbeitet:       {processed}")
    print(f"  Übersprungen:      {skipped}")
    print(f"  Konflikte (manuell): {conflicts}  "
          f"({conflicts / max(processed, 1) * 100:.1f}%)")
    print(f"  Fehler:            {errors}")
    print()
    print("  Tipp: Bei Unterbrechung mit --start-row N fortsetzen")
    print(f"        (N = letzte verarbeitete Zeile + 1)")
    print("=" * 65)


if __name__ == "__main__":
    main()
